import re
import csv
from io import BytesIO
from datetime import datetime, time
from flask import Blueprint, request, jsonify, current_app
from roles.notifs import trigger_notification

try:
    import xlrd
    from openpyxl import load_workbook
except ImportError:
    xlrd = None
    load_workbook = None

schedules_bp = Blueprint('schedules', __name__)
supabase = None  

def _normalize_date(val):
    if not val:
        return None
    val_str = str(val).strip().split(' ')[0]
    if re.match(r'^\d{4}-\d{2}-\d{2}$', val_str):
        return val_str
    for fmt in ('%m/%d/%Y', '%d/%m/%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(val_str, fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return val_str

def _normalize_time_str(val):
    if val is None or str(val).strip() == '':
        return None
    if isinstance(val, (datetime, time)):
        return val.strftime("%H:%M")
    if isinstance(val, float):
        if 0 <= val < 1.0:
            total_minutes = int(round(val * 24 * 60))
            hours = (total_minutes // 60) % 24
            minutes = total_minutes % 60
            return f"{hours:02d}:{minutes:02d}"
    val_str = str(val).strip()
    match = re.match(r'^(\d{1,2}):(\d{2})', val_str)
    if match:
        h, m = int(match.group(1)), int(match.group(2))
        return f"{h:02d}:{m:02d}"
    return val_str

def _normalize_xls_cell(value):
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value).strip()

def _notify_summary_saved(summary_id, staff_id, added=0, updated=0, editing=False):
    try:
        counts = []
        if added:
            counts.append(f'{added} {"trip" if added == 1 else "trips"} added')
        if updated:
            counts.append(f'{updated} {"trip" if updated == 1 else "trips"} updated')
        if not counts:
            return True
        notified = trigger_notification(
            title='Trip Summary Updated' if editing else 'Trip Summary Created',
            message=f'{summary_id}: {", ".join(counts)}.',
            target_user_id=staff_id,
            target_role='admin',
            source_tag='staff',
            db_client=supabase,
        )
        if not notified:
            current_app.logger.warning('Notification could not be saved for summary %s', summary_id)
    except Exception:
        current_app.logger.exception('Notification failed for summary %s', summary_id)

def _vehicle_type_only(raw_type):
    text = str(raw_type or '').strip()
    return re.sub(r'^\s*\d+\s*(seats?|seater)\s*[-–—:]?\s*', '', text, flags=re.IGNORECASE).strip()

def _is_missing_vehicle_type_column(error):
    code = str(getattr(error, 'code', ''))
    message = str(getattr(error, 'message', error))
    details = str(getattr(error, 'details', ''))
    return (
        (code == 'PGRST204' or 'PGRST204' in message)
        and 'vehicle_type' in f'{message} {details}'
    )

def _legacy_summary_payload(payload):
    """Map vehicle_type to bus_type until the database migration is applied."""
    legacy_payload = dict(payload)
    vehicle_type = legacy_payload.pop('vehicle_type', None)
    if vehicle_type and not legacy_payload.get('bus_type'):
        legacy_payload['bus_type'] = vehicle_type
    return legacy_payload

def _record_exists(table, id_field, value):
    if value in (None, ''):
        return True
    result = (
        supabase.table(table)
        .select(id_field)
        .eq(id_field, value)
        .limit(1)
        .execute()
    )
    return bool(result.data)

def _validate_trip_resources(payload):
    vehicle_id = payload.get('vehicle_id')
    driver_id = payload.get('driver_id')
    if vehicle_id is not None and not _record_exists('vehicle', 'vehicle_id', vehicle_id):
        return f"Vehicle ID {vehicle_id} does not exist."
    if driver_id is not None and not _record_exists('driver_profile', 'driver_id', driver_id):
        return f"Driver ID {driver_id} does not exist."
    return None

def _route_key(route_name):
    return re.sub(r'\s+', ' ', str(route_name or '').strip()).lower()

def _canonical_route_name(route_name):
    clean_name = re.sub(r'\s+', ' ', str(route_name or '').strip())
    return clean_name or 'Unspecified Route'

def _destination_payload(route_id=None, route_name=None):
    clean_name = _canonical_route_name(route_name)
    clean_key = _route_key(clean_name)
    if route_id:
        destination = (
            supabase.table('destination')
            .select('route_id, route_name')
            .eq('route_id', route_id)
            .maybe_single()
            .execute()
        )
        if destination.data:
            return {
                'route_id': destination.data.get('route_id'),
                'route_name': destination.data.get('route_name') or clean_name,
            }

    existing = (
        supabase.table('destination')
        .select('route_id, route_name')
        .ilike('route_name', clean_name)
        .limit(1)
        .execute()
    )
    if existing.data:
        row = existing.data[0] if isinstance(existing.data, list) else existing.data
        return {
            'route_id': row.get('route_id'),
            'route_name': row.get('route_name') or clean_name,
        }

    all_destinations = (
        supabase.table('destination')
        .select('route_id, route_name')
        .execute()
    )
    for destination in all_destinations.data or []:
        if _route_key(destination.get('route_name')) == clean_key:
            return {
                'route_id': destination.get('route_id'),
                'route_name': destination.get('route_name') or clean_name,
            }

    try:
        created = (
            supabase.table('destination')
            .insert({'route_name': clean_name, 'price': 0})
            .execute()
        )
        row = (created.data or [{}])[0]
    except Exception:
        existing_after_insert = (
            supabase.table('destination')
            .select('route_id, route_name')
            .ilike('route_name', clean_name)
            .limit(1)
            .execute()
        )
        row = (existing_after_insert.data or [{}])[0]

    return {
        'route_id': row.get('route_id'),
        'route_name': row.get('route_name') or clean_name,
    }

def _apply_destination(payload, route_id=None, route_name=None):
    destination = _destination_payload(route_id=route_id, route_name=route_name)
    payload['route_id'] = destination.get('route_id')
    payload['route_name'] = destination.get('route_name') or str(route_name or 'Unspecified Route').strip()
    return payload

def _insert_summary_rows(payloads):
    try:
        return supabase.table('trip_schedule').insert(payloads).execute()
    except Exception as error:
        if not _is_missing_vehicle_type_column(error):
            raise
        return supabase.table('trip_schedule').insert(
            [_legacy_summary_payload(row) for row in payloads]
        ).execute()

def _update_summary_row(trip_id, payload):
    try:
        return supabase.table('trip_schedule').update(payload).eq('trip_id', trip_id).execute()
    except Exception as error:
        if not _is_missing_vehicle_type_column(error):
            raise
        return supabase.table('trip_schedule').update(
            _legacy_summary_payload(payload)
        ).eq('trip_id', trip_id).execute()

def _to_int(value):
    if value in (None, ''):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None

def _to_float(value):
    if value in (None, ''):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None

def _trip_update_payload(data):
    passenger_count = _to_int(data.get('passenger_count'))
    seating_capacity = _to_int(data.get('seating_capacity'))

    utilization_rate = _to_float(data.get('utilization_rate'))
    if utilization_rate is None and passenger_count is not None and seating_capacity:
        utilization_rate = passenger_count / seating_capacity

    update_payload = {}
    text_fields = [
        'summary_id', 'schedule_date', 'working_day', 'bus_type', 'vehicle_type',
        'classification', 'ticket_no', 'departure_time',
        'estimated_arrival_time', 'remarks',
    ]
    for field in text_fields:
        if field in data:
            value = str(data.get(field) or '').strip() or None
            update_payload[field] = _vehicle_type_only(value) if field == 'bus_type' else value

    if 'vehicle_id' in data:
        update_payload['vehicle_id'] = _to_int(data.get('vehicle_id'))
    if 'driver_id' in data:
        update_payload['driver_id'] = _to_int(data.get('driver_id'))
    if 'company_id' in data:
        update_payload['company_id'] = _to_int(data.get('company_id'))
    if 'seating_capacity' in data:
        update_payload['seating_capacity'] = seating_capacity
    if 'passenger_count' in data:
        update_payload['passenger_count'] = passenger_count or 0
    if utilization_rate is not None:
        update_payload['utilization_rate'] = utilization_rate
    if 'route_id' in data or 'route_name' in data or 'route' in data:
        update_payload = _apply_destination(
            update_payload,
            route_id=_to_int(data.get('route_id')),
            route_name=data.get('route_name') or data.get('route'),
        )

    return update_payload

def _trip_payload_changed(existing, update_payload):
    for field, incoming in update_payload.items():
        current = existing.get(field)
        if field in ('vehicle_id', 'driver_id', 'company_id', 'seating_capacity', 'passenger_count', 'route_id'):
            if _to_int(current) != incoming:
                return True
        elif field == 'utilization_rate':
            current_float = _to_float(current)
            if current_float is None or incoming is None:
                if current_float != incoming:
                    return True
            elif abs(current_float - incoming) > 0.000001:
                return True
        else:
            current_text = str(current or '').strip() or None
            incoming_text = str(incoming or '').strip() or None
            if field == 'bus_type':
                current_text = _vehicle_type_only(current_text)
                incoming_text = _vehicle_type_only(incoming_text)
            if current_text != incoming_text:
                return True
    return False

@schedules_bp.route('/api/schedules/upload-summary-xls', methods=['POST'])
def upload_summary_xls():
    try:
        if 'file' not in request.files:
            return jsonify({"success": False, "error": "No file uploaded."}), 400

        uploaded = request.files['file']
        file_bytes = uploaded.read()
        filename = uploaded.filename.lower()

        rows = []
        if filename.endswith('.csv'):
            decoded_file = file_bytes.decode('utf-8-sig').splitlines()
            reader = csv.reader(decoded_file)
            rows = [row for row in reader]
        else:
            try:
                workbook = xlrd.open_workbook(file_contents=file_bytes)
                sheet = workbook.sheet_by_index(0)
                rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
            except Exception:
                if load_workbook is None:
                    return jsonify({"success": False, "error": "Missing excel parsing libraries."}), 500
                workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True)
                sheet = workbook.active
                rows = [[cell.value for cell in row] for row in sheet.iter_rows()]

        header_idx = -1
        is_db_dump = False
        company_name_raw = None

        for i, row in enumerate(rows):
            row_str = " ".join([str(x) for x in row if x is not None])
            if 'customer:' in row_str.lower():
                cust_part = row_str.split(':', 1)[-1].strip()
                company_name_raw = re.split(r'[\t\n,]', cust_part)[0].strip()

            str_row = [str(x).lower().strip() for x in row if x is not None]
            if 'trip_id' in str_row and 'schedule_date' in str_row:
                header_idx = i
                is_db_dump = True
                break
            elif any('route' in c for c in str_row) and any('driver' in c for c in str_row):
                header_idx = i
                break
        
        if header_idx == -1:
            return jsonify({"success": False, "error": "Could not locate valid header row."}), 400

        headers = [str(x).strip().lower() for x in rows[header_idx]]
        
        db_drivers = supabase.table('driver_profile').select('driver_id, full_name').execute().data or []
        db_vehicles = supabase.table('vehicle').select('vehicle_id, plate_number, bus_type, vehicle_type').execute().data or []
        db_companies = supabase.table('client_company').select('company_id, company_name').execute().data or []

        matched_company_id = None
        matched_company_name = None
        if company_name_raw:
            for c in db_companies:
                if c['company_name'].lower() in company_name_raw.lower() or company_name_raw.lower() in c['company_name'].lower():
                    matched_company_id = c['company_id']
                    matched_company_name = c['company_name']
                    break

        driver_id_to_name = {str(d['driver_id']): d['full_name'] for d in db_drivers}
        vehicle_id_to_plate = {str(v['vehicle_id']): v['plate_number'] for v in db_vehicles}
        vehicle_id_to_type = {str(v['vehicle_id']): v.get('vehicle_type') for v in db_vehicles}

        def clean_plate(p):
            return re.sub(r'[^A-Z0-9]', '', str(p or '').upper())
        vehicle_lookup = {clean_plate(v['plate_number']): v for v in db_vehicles}

        data_rows = []
        for row in rows[header_idx + 1:]:
            if not any(row): continue
            record = {}
            for idx, h in enumerate(headers):
                if idx >= len(row): continue
                val = row[idx]
                if val is None: val = ""
                
                if is_db_dump:
                    if h == 'schedule_date': record['schedule_date'] = _normalize_date(val)
                    elif h == 'working_day': record['working_day'] = _normalize_xls_cell(val).upper()
                    elif h == 'bus_type': record['bus_type'] = _normalize_xls_cell(val)
                    elif 'vehicle type' in h: record['vehicle_type'] = _normalize_xls_cell(val)
                    elif h == 'classification': record['classification'] = _normalize_xls_cell(val)
                    elif h == 'vehicle_id': record['vehicle_id'] = _normalize_xls_cell(val)
                    elif h == 'seating_capacity': 
                        try: record['seating_capacity'] = int(float(_normalize_xls_cell(val)))
                        except: record['seating_capacity'] = 14
                    elif h == 'ticket_no': record['ticket_no'] = _normalize_xls_cell(val)
                    elif h == 'driver_id': record['driver_id'] = _normalize_xls_cell(val)
                    elif h == 'route_name': record['route_name'] = _normalize_xls_cell(val)
                    elif h == 'passenger_count': 
                        try: record['passenger_count'] = int(float(_normalize_xls_cell(val)))
                        except: record['passenger_count'] = 0
                    elif h == 'departure_time': record['departure_time'] = _normalize_time_str(val)
                    elif h == 'estimated_arrival_time': record['estimated_arrival_time'] = _normalize_time_str(val)
                    elif h == 'utilization_rate': 
                        try: record['utilization_rate'] = float(_normalize_xls_cell(val))
                        except: record['utilization_rate'] = None
                    elif h == 'remarks': record['remarks'] = _normalize_xls_cell(val)
                else:
                    if 'date' in h and 'cr' not in h: record['schedule_date'] = _normalize_date(val)
                    elif 'working' in h: record['working_day'] = _normalize_xls_cell(val).upper()
                    elif 'vehicle type' in h: record['vehicle_type'] = _normalize_xls_cell(val)
                    elif 'type' in h: record['vehicle_type'] = _normalize_xls_cell(val)
                    elif 'class' in h: record['classification'] = _normalize_xls_cell(val)
                    elif 'plate' in h or 'bus no' in h: record['plate_number'] = _normalize_xls_cell(val)
                    elif 'capacity' in h or 'seating' in h:
                        try: record['seating_capacity'] = int(float(_normalize_xls_cell(val)))
                        except: record['seating_capacity'] = 14
                    elif 'ticket' in h: record['ticket_no'] = _normalize_xls_cell(val)
                    elif 'driver' in h: record['driver_name'] = _normalize_xls_cell(val)
                    elif 'route' in h: record['route_name'] = _normalize_xls_cell(val)
                    elif 'pass' in h or 'pax' in h:
                        try: record['passenger_count'] = int(float(_normalize_xls_cell(val)))
                        except: record['passenger_count'] = 0
                    elif 'dept' in h or 'dep' in h: record['departure_time'] = _normalize_time_str(val)
                    elif 'arriv' in h: record['estimated_arrival_time'] = _normalize_time_str(val)
                    elif 'util' in h:
                        raw_u = str(val).replace('%', '').strip()
                        try:
                            u_val = float(raw_u)
                            record['utilization_rate'] = u_val / 100.0 if u_val > 1.0 else u_val
                        except: record['utilization_rate'] = None
                    elif 'remark' in h: record['remarks'] = _normalize_xls_cell(val)

            if not record.get('route_name'): continue

            if is_db_dump:
                v_id = str(record.get('vehicle_id') or '').strip()
                d_id = str(record.get('driver_id') or '').strip()
                record['plate_number'] = vehicle_id_to_plate.get(v_id, 'Unassigned')
                if not record.get('vehicle_type'):
                    record['vehicle_type'] = vehicle_id_to_type.get(v_id, '')
                record['driver_name'] = driver_id_to_name.get(d_id, 'Unassigned')
            else:
                record['company_id'] = matched_company_id
                record['client_company'] = matched_company_name
                raw_plate = clean_plate(record.get('plate_number'))
                matched_veh = vehicle_lookup.get(raw_plate)
                if matched_veh:
                    record['vehicle_id'] = matched_veh['vehicle_id']
                    record['plate_number'] = matched_veh['plate_number']
                    if not record.get('vehicle_type'):
                        record['vehicle_type'] = matched_veh.get('vehicle_type')
                else:
                    record['vehicle_id'] = None

                raw_driver = str(record.get('driver_name') or '').strip().upper()
                matched_driver_id = None
                matched_driver_name = record.get('driver_name')
                parts = re.split(r'[\s.]+', raw_driver)
                parts = [p for p in parts if p]
                if len(parts) >= 2:
                    initial = parts[0][0]
                    surname = parts[-1]
                    for d in db_drivers:
                        d_name_upper = d['full_name'].upper()
                        if surname in d_name_upper and (d_name_upper.startswith(initial) or f" {initial}" in d_name_upper):
                            matched_driver_id = d['driver_id']
                            matched_driver_name = d['full_name']
                            break

                record['driver_id'] = matched_driver_id
                record['driver_name'] = matched_driver_name

            dep_time = record.get('departure_time') or '08:00'
            arr_time = record.get('estimated_arrival_time')
            if not arr_time or str(arr_time).strip() == '':
                try:
                    h, m = [int(x) for x in dep_time.split(':')[:2]]
                    arr_time = f"{(h + 1) % 24:02d}:{m:02d}"
                except Exception:
                    arr_time = '09:00'
            record['estimated_arrival_time'] = arr_time

            data_rows.append(record)

        return jsonify({"success": True, "rows": data_rows}), 200
    except Exception as exc:
        return jsonify({"success": False, "error": f"Failed to parse file: {str(exc)}"}), 500

@schedules_bp.route('/api/schedules/staff-options', methods=['GET'])
def get_staff_options():
    try:
        query = supabase.table('user_account').select('staff_id, full_name').eq('role', 'staff').execute()
        return jsonify({"success": True, "data": query.data}), 200
    except Exception as e:
        current_app.logger.exception('Fetch staff options failed')
        return jsonify({"success": False, "message": str(e)}), 500

@schedules_bp.route('/api/schedules/dispatch-options', methods=['GET'])
def get_dispatch_options():
    try:
        all_vehicles = supabase.table('vehicle').select('*').execute()
        raw_drivers = supabase.table('driver_profile').select(
            'driver_id, full_name, birthday, phone_no, date_hired, '
            'employment_status, is_backup, ml_classification'
        ).execute()

        return jsonify({
            "success": True,
            "vehicles": all_vehicles.data or [],
            "drivers": raw_drivers.data or [],
            "blocked": False
        }), 200
    except Exception as e:
        current_app.logger.exception('Fetch dispatch options failed')
        return jsonify({"success": False, "message": str(e)}), 500

@schedules_bp.route('/api/schedules/driver/<string:driver_id>', methods=['GET'])
def get_driver_trips(driver_id):
    try:
        query = supabase.table('trip_schedule').select(
            'trip_id, summary_id, schedule_date, working_day, route_id, route_name, '
            'vehicle_id, driver_id, bus_type, vehicle_type, classification, '
            'ticket_no, seating_capacity, passenger_count, departure_time, '
            'estimated_arrival_time, utilization_rate, remarks, trip_status'
        ).eq('driver_id', driver_id).order('schedule_date', desc=True).execute()
        return jsonify({"success": True, "data": query.data}), 200
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

@schedules_bp.route('/api/schedules/all', methods=['GET'])
def get_all_trips():
    try:
        trips = supabase.table('trip_schedule').select('*').order(
            'schedule_date', desc=True
        ).execute()
        vehicles = supabase.table('vehicle').select(
            'vehicle_id, plate_number, bus_type'
        ).execute()
        drivers = supabase.table('driver_profile').select(
            'driver_id, full_name'
        ).execute()

        v_map = {v['vehicle_id']: v['plate_number'] for v in vehicles.data if v.get('vehicle_id') is not None}
        d_map = {d['driver_id']: d['full_name'] for d in drivers.data if d.get('driver_id') is not None}

        formatted_trips = []
        for t in trips.data:
            t['plate_number'] = v_map.get(t.get('vehicle_id'), 'No Plate Assigned')
            t['driver_name'] = d_map.get(t.get('driver_id'), 'Unassigned')
            formatted_trips.append(t)

        return jsonify({"success": True, "data": formatted_trips}), 200
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

@schedules_bp.route('/api/schedules/staff-summary/<string:staff_uuid>', methods=['GET'])
def get_staff_trip_summary(staff_uuid):
    try:
        query = supabase.table('trip_schedule').select('*')
        trips = query.order('schedule_date', desc=False).execute()

        vehicles = supabase.table('vehicle').select(
            'vehicle_id, plate_number, bus_type'
        ).execute()
        drivers = supabase.table('driver_profile').select(
            'driver_id, full_name'
        ).execute()
        companies = supabase.table('client_company').select(
            'company_id, company_name, address'
        ).execute()

        v_map = {v['vehicle_id']: v for v in (vehicles.data or []) if v.get('vehicle_id') is not None}
        d_map = {d['driver_id']: d['full_name'] for d in (drivers.data or []) if d.get('driver_id') is not None}
        c_map = {c['company_id']: c['company_name'] for c in (companies.data or []) if c.get('company_id') is not None}
        company_rows = {c['company_id']: c for c in (companies.data or []) if c.get('company_id') is not None}

        formatted_trips = []
        for trip in trips.data or []:
            vehicle = v_map.get(trip.get('vehicle_id'), {})
            schedule_date = trip.get('schedule_date')
            working_day = trip.get('working_day')
            if not working_day and schedule_date:
                try:
                    working_day = datetime.strptime(str(schedule_date)[:10], '%Y-%m-%d').strftime('%A').upper()
                except Exception:
                    working_day = None

            seating_capacity = trip.get('seating_capacity')
            passenger_count = trip.get('passenger_count')
            utilization_rate = trip.get('utilization_rate')
            if utilization_rate is None and seating_capacity:
                try:
                    utilization_rate = float(passenger_count or 0) / float(seating_capacity)
                except Exception:
                    utilization_rate = None

            formatted_trips.append({
                **trip,
                'date': schedule_date,
                'working_day': working_day,
                'bus_type': trip.get('bus_type') or vehicle.get('bus_type'),
                'vehicle_type': (
                    trip.get('vehicle_type') or vehicle.get('vehicle_type')
                    or trip.get('bus_type') or vehicle.get('bus_type') or 'Unspecified'
                ),
                'plate_number': vehicle.get('plate_number', 'Unassigned'),
                'seating_capacity': seating_capacity,
                'driver_name': d_map.get(trip.get('driver_id'), 'Unassigned'),
                'client_company': c_map.get(trip.get('company_id'), 'Unknown Client'),
                'client_company_address': (
                    company_rows.get(trip.get('company_id'), {}).get('address')
                    or company_rows.get(trip.get('company_id'), {}).get('company_address')
                    or ''
                ),
                'utilization_rate': utilization_rate,
            })

        return jsonify({"success": True, "data": formatted_trips}), 200
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

@schedules_bp.route('/api/schedules/staff-summary', methods=['POST'])
def create_staff_trip_summary(data=None, notify=True):
    try:
        if data is None:
            data = request.get_json() or {}
        staff_id = data.get('staff_id')
        rows = data.get('rows')
        if not isinstance(rows, list) or not rows:
            rows = [data]

        if not staff_id:
            return jsonify({"success": False, "message": "Staff is required."}), 400

        def to_int(value):
            if value in (None, ''): return None
            try: return int(value)
            except (TypeError, ValueError): return None

        def to_float(value):
            if value in (None, ''): return None
            try: return float(value)
            except (TypeError, ValueError): return None

        summary_id = str(data.get('summary_id') or '').strip()
        if not summary_id:
            summary_id = f"SUM-{datetime.now().strftime('%Y%m%d-%H%M%S')}"

        payloads = []
        for row in rows:
            if not isinstance(row, dict):
                continue
            schedule_date = row.get('schedule_date') or row.get('date') or data.get('schedule_date') or data.get('date')
            if not schedule_date:
                continue

            passenger_count = to_int(row.get('passenger_count')) or 0
            seating_capacity = to_int(row.get('seating_capacity'))
            if seating_capacity and passenger_count > seating_capacity:
                return jsonify({"success": False, "message": "Passenger count cannot exceed seating capacity."}), 400

            utilization_rate = to_float(row.get('utilization_rate'))
            if utilization_rate is None and seating_capacity:
                utilization_rate = passenger_count / seating_capacity

            payload = {
                "summary_id": summary_id,
                "staff_id": staff_id,
                "company_id": to_int(row.get('company_id') or data.get('company_id')),
                "schedule_date": schedule_date,
                "working_day": row.get('working_day'),
                "bus_type": _vehicle_type_only(row.get('bus_type') or row.get('vehicle_type')),
                "vehicle_type": str(row.get('vehicle_type') or '').strip() or None,
                "classification": row.get('classification'),
                "vehicle_id": to_int(row.get('vehicle_id')),
                "seating_capacity": seating_capacity,
                "ticket_no": str(row.get('ticket_no') or '').strip() or None,
                "driver_id": to_int(row.get('driver_id')),
                "passenger_count": passenger_count,
                "departure_time": row.get('departure_time'),
                "estimated_arrival_time": row.get('estimated_arrival_time') or row.get('arrival_time'),
                "utilization_rate": utilization_rate,
                "remarks": row.get('remarks')
            }
            _apply_destination(
                payload,
                route_id=to_int(row.get('route_id')),
                route_name=row.get('route_name') or row.get('route'),
            )
            resource_error = _validate_trip_resources(payload)
            if resource_error:
                return jsonify({"success": False, "message": resource_error}), 400
            payloads.append(payload)

        if not payloads:
            return jsonify({"success": False, "message": "At least one trip summary row with a date is required."}), 400

        response = _insert_summary_rows(payloads)
        if notify:
            _notify_summary_saved(summary_id, staff_id, added=len(payloads))
        return jsonify({
            "success": True,
            "message": "Trip summary added.",
            "summary_id": summary_id,
            "data": response.data or []
        }), 201
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

@schedules_bp.route('/api/schedules/staff-summary/<int:trip_id>', methods=['PUT'])
def update_staff_trip_summary(trip_id, data=None, notify=True):
    try:
        if data is None:
            data = request.get_json() or {}

        def to_int(value):
            if value in (None, ''): return None
            try: return int(value)
            except (TypeError, ValueError): return None

        def to_float(value):
            if value in (None, ''): return None
            try: return float(value)
            except (TypeError, ValueError): return None

        passenger_count = to_int(data.get('passenger_count'))
        seating_capacity = to_int(data.get('seating_capacity'))
        if seating_capacity and passenger_count is not None and passenger_count > seating_capacity:
            return jsonify({"success": False, "message": "Passenger count cannot exceed seating capacity."}), 400

        utilization_rate = to_float(data.get('utilization_rate'))
        if utilization_rate is None and passenger_count is not None and seating_capacity:
            utilization_rate = passenger_count / seating_capacity

        update_payload = {}
        text_fields = [
            'summary_id', 'schedule_date', 'working_day', 'bus_type', 'vehicle_type',
            'classification', 'ticket_no', 'route_name', 'departure_time',
            'estimated_arrival_time', 'remarks',
        ]
        for field in text_fields:
            if field in data:
                value = str(data.get(field) or '').strip() or None
                update_payload[field] = _vehicle_type_only(value) if field == 'bus_type' else value

        if 'vehicle_id' in data:
            update_payload['vehicle_id'] = to_int(data.get('vehicle_id'))
        if 'driver_id' in data:
            update_payload['driver_id'] = to_int(data.get('driver_id'))
        if 'company_id' in data:
            update_payload['company_id'] = to_int(data.get('company_id'))
        if 'seating_capacity' in data:
            update_payload['seating_capacity'] = seating_capacity
        if 'passenger_count' in data:
            update_payload['passenger_count'] = passenger_count or 0
        if utilization_rate is not None:
            update_payload['utilization_rate'] = utilization_rate

        if 'route_id' in data or 'route_name' in data or 'route' in data:
            _apply_destination(
                update_payload,
                route_id=to_int(data.get('route_id')),
                route_name=data.get('route_name') or data.get('route'),
            )

        resource_error = _validate_trip_resources(update_payload)
        if resource_error:
            return jsonify({"success": False, "message": resource_error}), 400

        if not update_payload:
            return jsonify({"success": False, "message": "No trip summary fields to update."}), 400

        response = _update_summary_row(trip_id, update_payload)
        updated_rows = response.data or []
        if not updated_rows:
            return jsonify({"success": False, "message": "Trip summary row not found."}), 404

        updated_trip = updated_rows[0]
        if notify:
            _notify_summary_saved(
                updated_trip.get('summary_id') or f'TRIP-{trip_id}',
                updated_trip.get('staff_id'), updated=1, editing=True,
            )

        return jsonify({"success": True, "message": "Trip summary updated.", "data": updated_rows}), 200
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

@schedules_bp.route('/api/schedules/staff-summary/save', methods=['POST'])
def save_staff_trip_summary():
    """Save an editor submission and notify once, including mixed adds/edits."""
    try:
        data = request.get_json() or {}
        rows = data.get('rows')
        summary_id = str(data.get('summary_id') or '').strip()
        if not summary_id or not isinstance(rows, list) or not rows:
            return jsonify({'success': False, 'message': 'Summary ID and rows are required.'}), 400
        if any(not isinstance(row, dict) for row in rows):
            return jsonify({'success': False, 'message': 'Invalid summary row.'}), 400

        existing = supabase.table('trip_schedule').select('*').eq('summary_id', summary_id).execute().data or []
        if not existing:
            return jsonify({'success': False, 'message': 'Summary not found.'}), 404
        staff_id = existing[0].get('staff_id')
        existing_by_id = {int(row['trip_id']): row for row in existing}
        valid_ids = set(existing_by_id)
        updates, additions = [], []
        seen_ids = set()
        for row in rows:
            passengers = int(row.get('passenger_count') or 0)
            capacity = int(row.get('seating_capacity') or 0)
            if capacity and passengers > capacity:
                return jsonify({'success': False, 'message': 'Passenger count cannot exceed seating capacity.'}), 400
            if row.get('trip_id') is not None:
                trip_id = int(row['trip_id'])
                if trip_id not in valid_ids or trip_id in seen_ids:
                    return jsonify({'success': False, 'message': 'Invalid or duplicate trip in summary.'}), 400
                seen_ids.add(trip_id)
                row_payload = {**row, 'summary_id': summary_id}
                update_payload = _trip_update_payload(row_payload)
                if _trip_payload_changed(existing_by_id[trip_id], update_payload):
                    updates.append((trip_id, row_payload))
            else:
                if not (row.get('schedule_date') or row.get('date') or data.get('schedule_date')):
                    return jsonify({'success': False, 'message': 'A date is required for each new trip.'}), 400
                additions.append(row)

        for trip_id, row in updates:
            response, status = update_staff_trip_summary(trip_id, data=row, notify=False)
            if status != 200:
                return response, status
        if additions:
            response, status = create_staff_trip_summary(
                data={**data, 'staff_id': staff_id, 'rows': additions}, notify=False,
            )
            if status != 201:
                return response, status

        _notify_summary_saved(summary_id, staff_id, added=len(additions), updated=len(updates), editing=True)
        return jsonify({'success': True, 'summary_id': summary_id, 'added': len(additions), 'updated': len(updates)}), 200
    except (TypeError, ValueError):
        return jsonify({'success': False, 'message': 'Invalid trip counts or IDs.'}), 400
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500
